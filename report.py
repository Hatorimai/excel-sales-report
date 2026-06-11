"""
売上集計レポート生成スクリプト
---------------------------------
input/ フォルダ内の全売上Excelファイルを読み込み、
「支店別」「カテゴリ別」「月別推移」の3つの集計表を
書式付きのExcelレポートとして output/ に出力する。
"""

from pathlib import Path
from openpyxl.chart import LineChart, Reference
import pandas as pd
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

INPUT_DIR = Path("input")
OUTPUT_DIR = Path("output")
OUTPUT_FILE = OUTPUT_DIR / "売上集計レポート.xlsx"

# === 書式の定義(色や線をここに集約し、あとから一括変更できるようにする) ===
HEADER_FILL = PatternFill("solid", start_color="305496")   # 見出し行:濃い青
HEADER_FONT = Font(color="FFFFFF", bold=True)              # 見出し文字:白・太字
TOTAL_FILL = PatternFill("solid", start_color="D9E1F2")    # 合計行:薄い青
TOTAL_FONT = Font(bold=True)
THIN_BORDER = Border(*[Side(style="thin")] * 4)            # 上下左右すべて細罫線


def load_all_files(input_dir: Path) -> pd.DataFrame:
    """input/ 内の全xlsxを読み込み、1つのDataFrameに結合して返す"""
    # Excelがファイルを開いている間に作る一時ファイル(~$で始まる)を除外する。
    # 発注者がExcelを開いたままツールを実行しても落ちないようにするための防御
    files = sorted(
        f for f in input_dir.glob("*.xlsx")
        if not f.name.startswith("~$")
    )
    if not files:
        raise FileNotFoundError(
            f"{input_dir}/ にExcelファイルが見つかりません。"
            "先に generate_sample_data.py を実行してください。"
        )

    dfs = []
    for f in files:
        df = pd.read_excel(f)
        dfs.append(df)
        print(f"読み込み: {f.name}({len(df)}行)")

    combined = pd.concat(dfs, ignore_index=True)
    print(f"→ 結合完了: 全{len(combined)}行\n")
    return combined


def add_month_column(df: pd.DataFrame) -> pd.DataFrame:
    """日付列(例: 2026-01-15)から月列(例: 2026-01)を作る"""
    df = df.copy()
    df["月"] = df["日付"].str[:7]
    return df


def summarize(df: pd.DataFrame) -> dict[str, pd.DataFrame]:
    """3種類の集計表を作り、シート名→DataFrameの辞書で返す"""
    by_branch = df.groupby("支店").agg(
        売上合計=("売上金額", "sum"),
        販売数量=("数量", "sum"),
        取引件数=("売上金額", "count"),
    ).reset_index()

    by_category = df.groupby("カテゴリ").agg(
        売上合計=("売上金額", "sum"),
        販売数量=("数量", "sum"),
        取引件数=("売上金額", "count"),
    ).reset_index().sort_values("売上合計", ascending=False)

    monthly = pd.pivot_table(
        df, index="月", columns="支店", values="売上金額", aggfunc="sum",
    )
    monthly["全社合計"] = monthly.sum(axis=1)
    monthly = monthly.reset_index()

    return {
        "支店別集計": by_branch,
        "カテゴリ別集計": by_category,
        "月別推移": monthly,
    }


def add_total_row(ws):
    """シート末尾に「合計」行をSUM数式で追加する。

    Pythonで計算した数値を書き込むのではなく、Excelの数式(=SUM)を
    埋め込むのがポイント。発注者が後からデータを手修正しても
    合計が自動で再計算される「生きたExcel」として納品できる。
    """
    last_data_row = ws.max_row          # 現在の最終データ行
    total_row = last_data_row + 1       # 合計行はその次

    ws.cell(row=total_row, column=1, value="合計")
    for col in range(2, ws.max_column + 1):
        letter = get_column_letter(col)
        # 例: =SUM(B2:B4) のような数式を文字列で書き込む
        ws.cell(
            row=total_row, column=col,
            value=f"=SUM({letter}2:{letter}{last_data_row})",
        )


def style_sheet(ws):
    """1シート分の書式(見出し色・罫線・桁区切り・列幅)を整える"""
    # 1. 見出し行(1行目):青背景・白太字・中央揃え
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center")

    # 2. 全セルに罫線、数値セルに桁区切り表示(#,##0)を適用
    for row in ws.iter_rows():
        for cell in row:
            cell.border = THIN_BORDER
            if isinstance(cell.value, (int, float)):
                cell.number_format = "#,##0"

    # 3. 合計行(最終行):薄青背景・太字・数式セルも桁区切り
    for cell in ws[ws.max_row]:
        cell.fill = TOTAL_FILL
        cell.font = TOTAL_FONT
        cell.number_format = "#,##0"

    # 4. 列幅:列ごとの最長文字数に合わせて自動調整
    #    (全角文字は幅2としてカウントし、日本語見出しの切れを防ぐ)
    for col_cells in ws.columns:
        max_width = 0
        for cell in col_cells:
            if cell.value is not None:
                text = str(cell.value)
                width = sum(2 if ord(ch) > 255 else 1 for ch in text)
                max_width = max(max_width, width)
        letter = get_column_letter(col_cells[0].column)
        ws.column_dimensions[letter].width = max_width + 3  # 余白ぶん+3

    # 5. 見出し行を固定(スクロールしても1行目が常に見える)
    ws.freeze_panes = "A2"

def add_line_chart(ws):
    """月別推移シートに支店別の折れ線グラフを挿入する。

    グラフ範囲はセル番地を固定で書かず ws.max_row / max_column から
    計算する。月数や支店数が変わってもコードの修正なしで動くようにするため。
    """
    last_data_row = ws.max_row - 1   # 最終行は「合計」行なのでグラフから除外
    last_data_col = ws.max_column - 1  # 最終列は「全社合計」。支店との桁差で
                                       # 支店の動きが潰れて見えるため除外

    chart = LineChart()
    chart.title = "支店別 月別売上推移"
    chart.y_axis.title = "売上金額(円)"
    chart.x_axis.title = "月"
    chart.height = 10  # 単位はcm
    chart.width = 20

    # データ範囲:見出し行(1行目)を含めて指定し、titles_from_data=True で
    # 1行目を系列名(凡例の「大阪支店」など)として使う
    data = Reference(
        ws, min_col=2, max_col=last_data_col,
        min_row=1, max_row=last_data_row,
    )
    # 横軸ラベル:A列の月(見出しは含めない)
    categories = Reference(
        ws, min_col=1, min_row=2, max_row=last_data_row,
    )
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(categories)

    # 表の右隣(G2セルの位置)に配置
    ws.add_chart(chart, "G2")

def main():
    df = load_all_files(INPUT_DIR)
    df = add_month_column(df)
    sheets = summarize(df)

    OUTPUT_DIR.mkdir(exist_ok=True)
    try:
        with pd.ExcelWriter(OUTPUT_FILE, engine="openpyxl") as writer:
            for sheet_name, sheet_df in sheets.items():
                sheet_df.to_excel(writer, sheet_name=sheet_name, index=False)

            # pandasの書き込み後、同じファイルをopenpyxlとして触って書式を載せる
            for sheet_name in sheets:
                ws = writer.book[sheet_name]
                add_total_row(ws)
                style_sheet(ws)
                if sheet_name == "月別推移":
                    add_line_chart(ws)
                    print(f"書式適用+グラフ挿入: {sheet_name}")
                else:
                    print(f"書式適用: {sheet_name}")
    except PermissionError:
        # 出力先のファイルをExcelで開いたまま実行すると書き込めない。
        # 英語のトレースバックではなく「次に何をすべきか」を日本語で伝える
        print(f"エラー: {OUTPUT_FILE} が開かれているため書き込めません。")
        print("Excelでレポートを閉じてから、もう一度実行してください。")
        return  # ここで処理を終了(完了メッセージを出さないため)

    print(f"\n完了: {OUTPUT_FILE} を出力しました")


if __name__ == "__main__":
    main()